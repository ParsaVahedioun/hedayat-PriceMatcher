"""Export of the result table to Excel (openpyxl) and PDF (reportlab).

The Excel output deliberately mirrors the warehouse's own file byte-for-byte
in look: same three columns (کد کالا / نام کالا / موجودي), same fonts, same
header colour, same column widths and row heights - with a single "قیمت"
column appended at the end. Nothing about match quality (score, matched PDF
code/page, match type, status, note) is written out; that machinery lives
only in the in-app result table for filtering and is never exported.

The PDF export needs two things that are easy to get wrong and that produce
the classic "مربع‌های به‌هم‌ریخته" output when missed:

  * a font that actually contains Persian glyphs - Helvetica does not, and
    reportlab will silently draw nothing usable
  * shaping and bidi reordering - reportlab draws the code points in the
    order they are given, so the text has to be shaped and reversed first
"""
from __future__ import annotations

import os
from typing import Sequence

from core.errors import ExportError
from utils.config import load_config
from utils.logger import get_logger
from utils.number_utils import format_price

log = get_logger("exporter")

# ------------------------------------------------------------- excel style
# Copied from the warehouse's own "لیست اسامی کالا" file so the exported
# result opens looking like the exact same file, just with a price column
# added - same fonts, same header colour, same column widths/row heights.
HEADER_FONT_NAME = "B Mitra"
HEADER_FONT_SIZE = 15.75
DATA_FONT_NAME = "B Yekan"
DATA_FONT_SIZE = 13
HEADER_FILL_COLOR = "FF40E201"
HEADER_ROW_HEIGHT = 24.75
DATA_ROW_HEIGHT = 20.25

# (result key, header label, column width, number_format)
COLUMNS = [
    ("code", "کد کالا", 11.125, "@"),
    ("name", "نام کالا", 69.375, "General"),
    ("stock", "موجودي", 9.25, "##.###"),
    ("price", "قیمت", 16, "#,##0"),
]


def _code_value(value):
    """Store as a number when it is one, exactly like the source file."""
    if value is None:
        return None
    text = str(value).strip()
    if text.isdigit():
        try:
            return int(text)
        except ValueError:
            pass
    return text


def export_excel(rows: Sequence[dict], path: str) -> str:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover
        raise ExportError("کتابخانه openpyxl نصب نیست.", str(exc)) from exc

    try:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Price Match"
        sheet.sheet_view.rightToLeft = True

        header_font = Font(name=HEADER_FONT_NAME, size=HEADER_FONT_SIZE, bold=True)
        header_fill = PatternFill("solid", fgColor=HEADER_FILL_COLOR)
        header_align = Alignment(horizontal="center")
        data_font = Font(name=DATA_FONT_NAME, size=DATA_FONT_SIZE)

        for column, (_, label, _, _) in enumerate(COLUMNS, start=1):
            cell = sheet.cell(row=1, column=column, value=label)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
        sheet.row_dimensions[1].height = HEADER_ROW_HEIGHT

        for index, row in enumerate(rows, start=2):
            for column, (key, _, _, number_format) in enumerate(COLUMNS, start=1):
                value = _code_value(row.get(key)) if key == "code" else row.get(key)
                cell = sheet.cell(row=index, column=column, value=value)
                cell.font = data_font
                if number_format != "General":
                    cell.number_format = number_format
            sheet.row_dimensions[index].height = DATA_ROW_HEIGHT

        for column, (_, _, width, _) in enumerate(COLUMNS, start=1):
            sheet.column_dimensions[get_column_letter(column)].width = width

        _ensure_dir(path)
        workbook.save(path)
        workbook.close()
    except ExportError:
        raise
    except Exception as exc:
        log.exception("excel export failed")
        raise ExportError("ذخیره فایل Excel انجام نشد.", str(exc)) from exc
    log.info("excel exported: %s (%d rows)", path, len(rows))
    return path


# ------------------------------------------------------------------- pdf
def _rtl(text) -> str:
    """Shape + reorder Persian text so reportlab draws it correctly."""
    if text is None or text == "":
        return ""
    text = str(text)
    try:
        import arabic_reshaper
        from bidi.algorithm import get_display

        return get_display(arabic_reshaper.reshape(text))
    except Exception:
        return text


def _register_font(cfg: dict) -> str:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    for candidate in cfg.get("pdf_font_candidates", []):
        if candidate and os.path.isfile(candidate):
            try:
                pdfmetrics.registerFont(TTFont("AppFont", candidate))
                return "AppFont"
            except Exception:
                continue
    log.warning("no unicode font found, falling back to Helvetica")
    return "Helvetica"


PDF_COLUMNS = [
    ("price", "قیمت (تومان)", 35),
    ("stock", "موجودي", 25),
    ("name", "نام کالا", 130),
    ("code", "کد کالا", 24),
]


def export_pdf(rows: Sequence[dict], path: str, stats: dict | None = None) -> str:
    cfg = load_config()
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError as exc:  # pragma: no cover
        raise ExportError("کتابخانه reportlab نصب نیست.", str(exc)) from exc

    try:
        font = _register_font(cfg)
        _ensure_dir(path)
        doc = SimpleDocTemplate(
            path, pagesize=landscape(A4),
            rightMargin=8 * mm, leftMargin=8 * mm,
            topMargin=10 * mm, bottomMargin=10 * mm,
            title="Price Match Report",
        )
        title_style = ParagraphStyle("title", fontName=font, fontSize=14, alignment=1, leading=18)
        info_style = ParagraphStyle("info", fontName=font, fontSize=10, alignment=1, leading=14)

        story = [Paragraph(_rtl("لیست قیمت"), title_style)]
        brand = (stats or {}).get("brand")
        if brand:
            story.append(Spacer(1, 2 * mm))
            story.append(Paragraph(_rtl(f"برند: {brand}"), info_style))
        story.append(Spacer(1, 4 * mm))

        data = [[_rtl(label) for _, label, _ in PDF_COLUMNS]]
        for row in rows:
            line = []
            for key, _, _ in PDF_COLUMNS:
                if key == "price":
                    line.append(format_price(row.get("price")))
                elif key == "stock":
                    line.append(format_price(row.get("stock")))
                elif key == "name":
                    line.append(_rtl(str(row.get("name", ""))[:80]))
                else:
                    line.append(str(row.get(key) or "-"))
            data.append(line)

        table = Table(data, colWidths=[width * mm for _, _, width in PDF_COLUMNS],
                      repeatRows=1)
        style = [
            ("FONTNAME", (0, 0), (-1, -1), font),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#40E201")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.black),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#999999")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F6FB")]),
        ]
        table.setStyle(TableStyle(style))
        story.append(table)
        doc.build(story)
    except ExportError:
        raise
    except Exception as exc:
        log.exception("pdf export failed")
        raise ExportError("ساخت فایل PDF انجام نشد.", str(exc)) from exc
    log.info("pdf exported: %s (%d rows)", path, len(rows))
    return path


def _ensure_dir(path: str) -> None:
    folder = os.path.dirname(os.path.abspath(path))
    if folder:
        os.makedirs(folder, exist_ok=True)
