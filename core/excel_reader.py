"""Excel inventory reader.

openpyxl in `read_only` + `values_only` mode: rows are streamed, never loaded
as a whole workbook object graph, and no DataFrame copy is created.
pandas is deliberately not used here - it would add ~50MB of RAM and a heavy
import for something a plain loop does better.

The inventory is the source of truth for the output, so this reader is the
one that decides how many result rows there will be.
"""
from __future__ import annotations

import os
from typing import Iterable

from core.errors import (
    MSG_EXCEL_COLUMNS,
    MSG_EXCEL_EMPTY,
    MSG_EXCEL_UNREADABLE,
    ExcelError,
)
from core.normalizer import map_header
from models.product import StockItem
from utils.logger import get_logger
from utils.number_utils import parse_number

log = get_logger("excel")

HEADER_SCAN_ROWS = 20
REQUIRED_ANY = ("code", "name")


def _detect_header(rows: list[tuple]) -> tuple[int, dict[int, str]]:
    """Return (header_row_index, {column_index: field}) for the best row."""
    best_row, best_map, best_count = -1, {}, 0
    for row_idx, row in enumerate(rows):
        mapping: dict[int, str] = {}
        for col_idx, cell in enumerate(row):
            if cell is None:
                continue
            field = map_header(str(cell))
            if field and field not in mapping.values():
                mapping[col_idx] = field
        count = len(mapping)
        if count > best_count:
            best_row, best_map, best_count = row_idx, mapping, count
    if best_count < 2:
        return -1, {}
    return best_row, best_map


def read_excel(path: str, max_rows: int | None = None) -> list[StockItem]:
    """Read an inventory file into a list of StockItems."""
    if not path or not os.path.isfile(path):
        raise ExcelError(MSG_EXCEL_UNREADABLE, f"file not found: {path}")

    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ExcelError("کتابخانه openpyxl نصب نیست.", str(exc)) from exc

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        log.exception("cannot open excel")
        raise ExcelError(MSG_EXCEL_UNREADABLE, str(exc)) from exc

    try:
        sheet = _pick_sheet(workbook)
        head_buffer: list[tuple] = []
        row_iter = sheet.iter_rows(values_only=True)
        for row in row_iter:
            head_buffer.append(row)
            if len(head_buffer) >= HEADER_SCAN_ROWS:
                break

        header_row, mapping = _detect_header(head_buffer)
        if header_row < 0 or not any(f in mapping.values() for f in REQUIRED_ANY):
            raise ExcelError(MSG_EXCEL_COLUMNS, f"headers: {head_buffer[:3]}")
        if "stock" not in mapping.values():
            log.warning("no stock column detected")

        items: list[StockItem] = []
        data_rows: Iterable[tuple] = head_buffer[header_row + 1:]
        row_no = header_row + 1
        for chunk in (data_rows, row_iter):
            for row in chunk:
                row_no += 1
                item = _row_to_item(row, mapping, row_no)
                if item is not None:
                    items.append(item)
                if max_rows and len(items) >= max_rows:
                    break
            if max_rows and len(items) >= max_rows:
                break
    finally:
        try:
            workbook.close()
        except Exception:
            pass

    if not items:
        raise ExcelError(MSG_EXCEL_EMPTY, "no data rows")
    log.info("excel: %d rows, mapping=%s", len(items), mapping)
    return items


def _pick_sheet(workbook):
    """Use the active sheet unless another one clearly holds more data."""
    best = workbook.active
    best_size = (best.max_row or 0) * (best.max_column or 0)
    for sheet in workbook.worksheets:
        size = (sheet.max_row or 0) * (sheet.max_column or 0)
        if size > best_size * 2:
            best, best_size = sheet, size
    return best


def _row_to_item(row: tuple, mapping: dict[int, str], row_no: int) -> StockItem | None:
    if not row:
        return None
    item = StockItem(row=row_no)
    has_data = False
    for col_idx, field in mapping.items():
        if col_idx >= len(row):
            continue
        value = row[col_idx]
        if value is None or (isinstance(value, str) and not value.strip()):
            continue
        has_data = True
        if field == "stock":
            item.stock = parse_number(value)
        elif field == "code":
            item.code = str(value).strip()
        elif field == "name":
            item.name = str(value).strip()
        elif field == "unit":
            item.unit = str(value).strip()
    if not has_data or (not item.code and not item.name):
        return None
    return item
